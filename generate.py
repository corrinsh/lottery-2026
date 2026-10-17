# -*- coding: utf-8 -*-
"""生成两张数据表 + 两类折线图（纯数据记录，不含任何预测）。"""
import json, os
from datetime import datetime
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

BASE = r"E:\WB工作区\数据\彩票数据研究_2026"
CH = os.path.join(BASE, "charts")
os.makedirs(CH, exist_ok=True)

# ---------- 中文字体 ----------
for f in ["Microsoft YaHei", "SimHei", "DengXian"]:
    try:
        font_manager.findfont(f, fallback_to_default=False)
        plt.rcParams["font.sans-serif"] = [f]
        break
    except Exception:
        continue
plt.rcParams["axes.unicode_minus"] = False
plt.rcParams["figure.dpi"] = 130

dlt = json.load(open(os.path.join(BASE, "data", "dlt.json"), encoding="utf-8"))
ssq = json.load(open(os.path.join(BASE, "data", "ssq.json"), encoding="utf-8"))

DLT_COLS = ["期号", "开奖日期", "前区1", "前区2", "前区3", "前区4", "前区5",
            "后区1", "后区2", "前区和值", "前区跨度", "后区和值"]
SSQ_COLS = ["期号", "开奖日期", "红球1", "红球2", "红球3", "红球4", "红球5", "红球6",
            "蓝球", "红球和值", "红球跨度", "奇偶比(红球)"]


def enrich(rows, kind):
    for r in rows:
        if kind == "dlt":
            pre = [r[f"前区{i}"] for i in range(1, 6)]
            r["前区跨度"] = max(pre) - min(pre)
            r["后区和值"] = r["后区1"] + r["后区2"]
        else:
            red = [r[f"红球{i}"] for i in range(1, 7)]
            r["红球跨度"] = max(red) - min(red)
            r["奇偶比(红球)"] = f"{sum(1 for x in red if x % 2 == 1)}:{sum(1 for x in red if x % 2 == 0)}"
    return rows


enrich(dlt, "dlt")
enrich(ssq, "ssq")

# ---------- Excel ----------
wb = Workbook()
thin = Side(style="thin", color="D9D9D9")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
head_fill = PatternFill("solid", fgColor="1F4E79")
head_font = Font(bold=True, color="FFFFFF", size=11)
zebra = PatternFill("solid", fgColor="F2F7FB")


def build_sheet(ws, rows, cols, title, color_map=None):
    ws.append(cols)
    for c in range(1, len(cols) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    ws.row_dimensions[1].height = 26
    for i, r in enumerate(rows, start=2):
        ws.append([r[c] for c in cols])
        for c in range(1, len(cols) + 1):
            cell = ws.cell(row=i, column=c)
            cell.alignment = Alignment(horizontal="center")
            cell.border = border
            if i % 2 == 0:
                cell.fill = zebra
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(rows)+1}"
    widths = {1: 11, 2: 13}
    for idx, name in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = widths.get(idx, 11)
    ws.sheet_view.zoomScale = 100


ws1 = wb.active
ws1.title = "大乐透"
build_sheet(ws1, dlt, DLT_COLS, "大乐透")
ws2 = wb.create_sheet("双色球")
build_sheet(ws2, ssq, SSQ_COLS, "双色球")

# 说明页
ws0 = wb.create_sheet("数据说明", 0)
notes = [
    ("数据说明", ""),
    ("用途声明", "本表仅用于公开开奖数据的记录与可视化研究，不含任何预测、推荐或投注建议。"),
    ("数据来源（大乐透）", "中国体育彩票官方网站 · 超级大乐透历史开奖公告接口"),
    ("数据来源（双色球）", "中国福利彩票官方网站 · 双色球历史开奖公告接口"),
    ("抓取时间", datetime.now().strftime("%Y-%m-%d %H:%M") + " (GMT+8)"),
    ("统计区间", "2026-01-01 至 2026-10-18"),
    ("大乐透期数", f"{len(dlt)} 期（{dlt[0]['期号']} — {dlt[-1]['期号']}，{dlt[0]['开奖日期']} — {dlt[-1]['开奖日期']}）"),
    ("双色球期数", f"{len(ssq)} 期（{ssq[0]['期号']} — {ssq[-1]['期号']}，{ssq[0]['开奖日期']} — {ssq[-1]['开奖日期']}）"),
    ("休市说明", "区间内包含春节休市（2 月中旬）与国庆休市（10 月初），因此相邻期日期间隔会超过常规开奖间隔，非数据缺失。"),
    ("口径 · 大乐透", "前区 5 个号码（1-35）+ 后区 2 个号码（1-12），共 7 个号码；表内按开奖公告升序排列。"),
    ("口径 · 双色球", "红球 6 个号码（1-33）+ 蓝球 1 个号码（1-16），共 7 个号码；表内按开奖公告升序排列。"),
    ("和值定义", "该组号码的算术和，属描述统计量。"),
    ("跨度定义", "该组号码最大值减最小值。"),
    ("图表说明", "charts 目录：*_seq.png 为 7 个号码逐期连续折线；*_pos1~7.png 为第 1 至第 7 位各自的单期走势。"),
]
for r, (k, v) in enumerate(notes, start=1):
    ws0.cell(row=r, column=1, value=k)
    ws0.cell(row=r, column=2, value=v)
    if r == 1:
        ws0.cell(row=r, column=1).font = Font(bold=True, size=14, color="1F4E79")
    else:
        ws0.cell(row=r, column=1).font = Font(bold=True)
        ws0.cell(row=r, column=1).alignment = Alignment(vertical="top")
ws0.column_dimensions["A"].width = 22
ws0.column_dimensions["B"].width = 105
for r in range(2, len(notes) + 1):
    ws0.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    ws0.row_dimensions[r].height = 22

xlsx = os.path.join(BASE, "2026年大乐透与双色球开奖数据.xlsx")
wb.save(xlsx)

# 另附 CSV（UTF-8 BOM，Excel 双击不乱码）
import csv
for rows, cols, name in [(dlt, DLT_COLS, "大乐透"), (ssq, SSQ_COLS, "双色球")]:
    with open(os.path.join(BASE, f"{name}数据表_2026.csv"), "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(cols)
        for r in rows:
            w.writerow([r[c] for c in cols])

# ---------- 折线图 ----------
BG = "#FFFFFF"
INK = "#1A1A1A"
GRID = "#E5E7EB"


def style_ax(ax, ylabel, xlabel):
    ax.set_facecolor(BG)
    ax.grid(True, color=GRID, linewidth=0.8, linestyle="-")
    ax.set_axisbelow(True)
    for s in ["top", "right"]:
        ax.spines[s].set_visible(False)
    for s in ["left", "bottom"]:
        ax.spines[s].set_color("#9CA3AF")
    ax.set_ylabel(ylabel, fontsize=10, color=INK)
    ax.set_xlabel(xlabel, fontsize=10, color=INK)
    ax.tick_params(colors="#4B5563", labelsize=9)


POS_COLORS = ["#1D4ED8", "#0891B2", "#059669", "#CA8A04", "#DC2626", "#7C3AED", "#DB2777"]


def seq_chart(rows, kind, fname, title, npos=7):
    """所有期所有号码按顺序连成一条线：每期 7 个点，一期接一期；下附最近 20 期放大。"""
    labels = ["前区1", "前区2", "前区3", "前区4", "前区5", "后区1", "后区2"] if kind == "dlt" \
        else ["红球1", "红球2", "红球3", "红球4", "红球5", "红球6", "蓝球"]
    xs, ys, posidx, period_starts = [], [], [], []
    k = 0
    for r in rows:
        period_starts.append(k)
        for j, col in enumerate(labels):
            xs.append(k); ys.append(r[col]); posidx.append(j); k += 1

    fig = plt.figure(figsize=(16, 8.6), facecolor=BG)
    gs = fig.add_gridspec(2, 1, height_ratios=[2.5, 1.15], hspace=0.34)
    ax = fig.add_subplot(gs[0])
    ax2 = fig.add_subplot(gs[1])

    def draw(a, lo_x, hi_x, xlabels, xpos, rotate, tag_all):
        style_ax(a, "号码数值", "")
        a.plot(xs, ys, color="#B6C2D1", linewidth=0.55, alpha=0.9, zorder=1)
        for j, col in enumerate(labels):
            ii = [i for i, q in enumerate(posidx) if q == j]
            a.scatter([xs[i] for i in ii], [ys[i] for i in ii], s=7 if tag_all else 26,
                      color=POS_COLORS[j], zorder=3, linewidth=0,
                      label=f"第{j+1}位 · {col}")
        for s in period_starts:
            if lo_x <= s <= hi_x:
                a.axvline(s, color="#DCE3EA", linewidth=0.4, alpha=0.9, zorder=0)
        a.set_xlim(lo_x, hi_x)
        a.set_xticks(xpos)
        a.set_xticklabels(xlabels, rotation=rotate, ha="right", fontsize=8)

    tick_step = 10
    ti = list(range(0, len(rows), tick_step))
    draw(ax, -4, k + 3, [rows[i]["期号"] for i in ti],
         [period_starts[i] for i in ti], 60, True)
    ax.legend(loc="upper left", fontsize=8, frameon=False, ncol=7,
              handletextpad=0.2, columnspacing=0.9)
    ax.set_title(title, fontsize=14, color=INK, pad=30, fontweight="bold")

    # 放大：最近 20 期
    start = len(rows) - 20
    lo = period_starts[start] - 1
    draw(ax2, lo, k + 3, [rows[i]["期号"] for i in range(start, len(rows))],
         [period_starts[i] for i in range(start, len(rows))], 90, False)
    ax2.set_title(f"放大：最近 20 期（{rows[start]['期号']} 起）—— 可见每期 7 个点的完整锯齿",
                  fontsize=10.5, color="#374151", pad=8)
    ax2.set_xlabel("期号（每期 7 个号码，按第一位到第七位顺序衔接）", fontsize=10, color=INK)

    p = os.path.join(CH, fname)
    fig.savefig(p, facecolor=BG, bbox_inches="tight")
    plt.close(fig)
    return p


def pos_charts(rows, kind, prefix, pos_defs):
    """每个位置单独一张图：只画第 N 位的走势。"""
    out = []
    for idx, (col, label, lo, hi) in enumerate(pos_defs, start=1):
        vals = [r[col] for r in rows]
        mean = sum(vals) / len(vals)
        xs = list(range(1, len(vals) + 1))

        fig, ax = plt.subplots(figsize=(14, 4.6), facecolor=BG)
        style_ax(ax, f"{label}号码值", "期序号（从左到右为 2026 年 1 月 → 10 月）")
        ax.plot(xs, vals, color="#0F766E", linewidth=1.2, marker="o", markersize=3,
                markerfacecolor=BG, markeredgewidth=1.0, markeredgecolor="#0F766E",
                label=f"第{idx}位实际开奖号码")
        ax.axhline(mean, color="#DC2626", linewidth=1.1, linestyle="--",
                   label=f"全期算术平均 = {mean:.2f}（描述统计，非预测）")
        ax.fill_between(xs, mean, vals, color="#0F766E", alpha=0.07)

        tick_step = max(1, len(vals) // 14)
        ti = list(range(0, len(vals), tick_step))
        ax.set_xticks([xs[i] for i in ti])
        ax.set_xticklabels([rows[i]["期号"] for i in ti], rotation=60, ha="right", fontsize=8)
        ax.set_ylim(lo - 1.2, hi + 1.2)
        ax.set_xlim(0, len(vals) + 1)
        ax.legend(loc="upper left", fontsize=9, frameon=False)
        ax.set_title(f"{prefix} · 第{idx}位号码逐期走势（共 {len(vals)} 期）",
                     fontsize=13, color=INK, pad=12, fontweight="bold")
        fig.tight_layout()
        p = os.path.join(CH, f"{'dlt' if kind=='dlt' else 'ssq'}_pos{idx}.png")
        fig.savefig(p, facecolor=BG, bbox_inches="tight")
        plt.close(fig)
        out.append(p)
    return out


dlt_pos = [("前区1", "前区第1位", 1, 35), ("前区2", "前区第2位", 1, 35),
           ("前区3", "前区第3位", 1, 35), ("前区4", "前区第4位", 1, 35),
           ("前区5", "前区第5位", 1, 35), ("后区1", "后区第1位", 1, 12),
           ("后区2", "后区第2位", 1, 12)]
ssq_pos = [("红球1", "红球第1位", 1, 33), ("红球2", "红球第2位", 1, 33),
           ("红球3", "红球第3位", 1, 33), ("红球4", "红球第4位", 1, 33),
           ("红球5", "红球第5位", 1, 33), ("红球6", "红球第6位", 1, 33),
           ("蓝球", "蓝球（第7位）", 1, 16)]

files = []
files.append(seq_chart(dlt, "dlt", "dlt_seq.png", "大乐透 · 七个号码逐期连续折线（2026-01-03 至 2026-10-17，118 期）"))
files.append(seq_chart(ssq, "ssq", "ssq_seq.png", "双色球 · 七个号码逐期连续折线（2026-01-01 至 2026-10-15，118 期）"))
files += pos_charts(dlt, "dlt", "大乐透", dlt_pos)
files += pos_charts(ssq, "ssq", "双色球", ssq_pos)

# ---------- 统计摘要 ----------
def summary(rows, kind):
    s = {"期数": len(rows), "起始期": rows[0]["期号"], "截止期": rows[-1]["期号"],
         "起始日期": rows[0]["开奖日期"], "截止日期": rows[-1]["开奖日期"]}
    if kind == "dlt":
        pre = [r[f"前区{i}"] for r in rows for i in range(1, 6)]
        s["前区均值"] = round(sum(pre) / len(pre), 2)
        s["后区均值"] = round(sum(r["后区1"] + r["后区2"] for r in rows) / (2 * len(rows)), 2)
    else:
        red = [r[f"红球{i}"] for r in rows for i in range(1, 7)]
        s["红球均值"] = round(sum(red) / len(red), 2)
        s["蓝球均值"] = round(sum(r["蓝球"] for r in rows) / len(rows), 2)
    return s

print(json.dumps({"dlt": summary(dlt, "dlt"), "ssq": summary(ssq, "ssq")}, ensure_ascii=False, indent=1))
print("files:", len(files))

# ---------- HTML 图册 ----------
sd, ss = summary(dlt, "dlt"), summary(ssq, "ssq")


def blocks(prefix, title, meta, s):
    pos = "".join(
        f'<figure><img src="charts/{prefix}_pos{i}.png" alt="{title} 第{i}位走势"><figcaption>第 {i} 位号码逐期走势</figcaption></figure>'
        for i in range(1, 8))
    return f"""
  <section>
    <h2>{title}</h2>
    <p class="meta">{meta}</p>
    <div class="stats">
      <div><span>期数</span><b>{s['期数']}</b></div>
      <div><span>区间</span><b>{s['起始期']} — {s['截止期']}</b></div>
      <div><span>日期</span><b>{s['起始日期']} → {s['截止日期']}</b></div>
    </div>
    <h3>① 七个号码逐期连续折线</h3>
    <figure class="wide"><img src="charts/{prefix}_seq.png" alt="{title} 七位连续折线"><figcaption>每期 7 个号码按第 1 → 第 7 位顺序排列，一期接一期连续衔接</figcaption></figure>
    <h3>② 单个位次走势（第 1 — 第 7 位，共 7 张）</h3>
    <div class="grid">{pos}</div>
  </section>"""


html = f"""<!DOCTYPE html>
<html lang="zh-CN"><head><meta charset="utf-8"><title>2026 年大乐透 / 双色球开奖数据研究</title>
<style>
  :root{{--bg:#0f1216;--card:#171b21;--line:#262c35;--ink:#e8edf3;--dim:#8b97a6;--accent:#4da3ff;}}
  *{{box-sizing:border-box}}
  body{{margin:0;background:var(--bg);color:var(--ink);font:15px/1.7 "Microsoft YaHei",system-ui,sans-serif}}
  .wrap{{max-width:1240px;margin:0 auto;padding:40px 24px 80px}}
  h1{{font-size:26px;margin:0 0 6px}}
  .sub{{color:var(--dim);font-size:13.5px;margin-bottom:22px}}
  .notice{{background:#1b1410;border:1px solid #4a2f1c;border-left:3px solid #d97706;padding:12px 16px;border-radius:8px;font-size:13.5px;color:#f0d9be;margin-bottom:34px}}
  section{{margin-bottom:56px}}
  h2{{font-size:20px;margin:0 0 4px;padding-bottom:10px;border-bottom:1px solid var(--line)}}
  h3{{font-size:14px;color:var(--accent);margin:26px 0 12px;font-weight:600;letter-spacing:.3px}}
  .meta{{color:var(--dim);font-size:13px;margin:8px 0 14px}}
  .stats{{display:flex;gap:10px;flex-wrap:wrap;margin-bottom:10px}}
  .stats>div{{background:var(--card);border:1px solid var(--line);border-radius:8px;padding:8px 14px;font-size:12.5px}}
  .stats span{{color:var(--dim);margin-right:8px}}
  .stats b{{color:var(--ink);font-weight:600}}
  figure{{margin:0;background:#fff;border-radius:10px;overflow:hidden;border:1px solid var(--line)}}
  figure img{{display:block;width:100%;cursor:zoom-in}}
  figcaption{{background:var(--card);color:var(--dim);font-size:12px;padding:8px 12px;border-top:1px solid var(--line)}}
  .grid{{display:grid;grid-template-columns:1fr 1fr;gap:16px}}
  @media(max-width:900px){{.grid{{grid-template-columns:1fr}}}}
  #lb{{position:fixed;inset:0;background:rgba(0,0,0,.94);display:none;align-items:center;justify-content:center;z-index:99;cursor:zoom-out;padding:20px}}
  #lb img{{max-width:98%;max-height:98%}}
  footer{{color:var(--dim);font-size:12px;border-top:1px solid var(--line);padding-top:16px}}
  footer p{{margin:6px 0}}
</style></head><body><div class="wrap">
  <h1>2026 年大乐透 / 双色球开奖数据研究</h1>
  <div class="sub">数据来源：中国体育彩票、中国福利彩票官方网站开奖公告　|　区间：{dlt[0]['开奖日期']} — {dlt[-1]['开奖日期']}　|　数据截至 {dlt[-1]['期号']} 期</div>
  <div class="notice"><b>声明：</b>本页仅对已公布的历史开奖号码做记录与可视化，不含任何预测、推荐或投注建议。图中红色虚线为全期算术平均值，属描述统计量，不代表未来结果。</div>
  {blocks('dlt', '大乐透', '前区 5 个（1—35）+ 后区 2 个（1—12），每期共 7 个号码', sd)}
  {blocks('ssq', '双色球', '红球 6 个（1—33）+ 蓝球 1 个（1—16），每期共 7 个号码', ss)}
  <footer>
    <p>引用：中国体育彩票. “超级大乐透历史开奖公告.” <i>中国体育彩票</i>, 国家体育总局体育彩票管理中心, webapi.sporttery.cn. 2026 年 10 月 18 日访问.</p>
    <p>引用：中国福利彩票. “双色球历史开奖公告.” <i>中国福利彩票</i>, 中国福利彩票发行管理中心, www.cwl.gov.cn. 2026 年 10 月 18 日访问.</p>
    <p>注：区间内春节（2 月中旬）与国庆（10 月初）休市造成的日期空档不属于数据缺失。</p>
  </footer>
</div>
<div id="lb"><img src="" alt=""></div>
<script>
  const lb = document.getElementById('lb'), lbi = lb.querySelector('img');
  document.querySelectorAll('figure img').forEach(img => img.onclick = () => {{
    lbi.src = img.src; lb.style.display = 'flex';
  }});
  lb.onclick = () => lb.style.display = 'none';
</script></body></html>"""

with open(os.path.join(BASE, "开奖数据图表总览.html"), "w", encoding="utf-8") as f:
    f.write(html)
print("html saved")
